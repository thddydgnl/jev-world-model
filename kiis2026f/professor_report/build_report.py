from pathlib import Path
import sys, json, csv, math
sys.path.insert(0, '/private/tmp/kiis_professor_plotdeps')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
QA = BASE / '.qa'
QA.mkdir(exist_ok=True)
S = json.loads((ROOT/'results/slots.json').read_text())['slots']
FONT = 'Noto Sans CJK KR'
FONT_FILE = '/System/Library/Fonts/AppleSDGothicNeo.ttc'

def make_flow():
    im=Image.new('RGB',(1900,390),'white'); d=ImageDraw.Draw(im)
    f=ImageFont.truetype(FONT_FILE,38); small=ImageFont.truetype(FONT_FILE,29)
    labels=[('현재 상태','완전관측'),('정책','후보 행동열'),('월드 모델','전이 예측'),('계획기','효용 비교'),('실제 환경','첫 행동 실행')]
    for i,(a,b) in enumerate(labels):
        x=20+i*380
        d.rounded_rectangle((x,45,x+320,205),radius=12,fill='#f1f3f5',outline='#777777',width=2)
        d.text((x+160,99),a,font=f,fill='black',anchor='mm');d.text((x+160,157),b,font=small,fill='#333333',anchor='mm')
        if i<4:
            d.line((x+330,125,x+365,125),fill='#444444',width=4)
            d.polygon([(x+365,125),(x+350,116),(x+350,134)],fill='#444444')
    d.line([(1700,220),(1700,280),(180,280),(180,220)],fill='#555555',width=3)
    d.polygon([(180,220),(171,237),(189,237)],fill='#555555')
    d.text((950,326),'새 관측으로 갱신하고 다음 step에서 재계획',font=small,fill='#333333',anchor='mm')
    im.save(QA/'architecture.png')

def make_rollout():
    rows=list(csv.DictReader((ROOT/'results/fig2.csv').open()))
    colors={'A_jev':'#145f98','B0_typed':'#8894a0','B_typed':'#16806a','D0_gen':'#c69486','D_gen':'#b44a42','persistence':'#333333'}
    labels={'A_jev':'A  JEV','B0_typed':'B0  Decision','B_typed':'B  Decision + LoRA','D0_gen':'D0  Generative','D_gen':'D  Generative + LoRA','persistence':'Persistence'}
    plt.rcParams.update({'font.size':9,'font.family':'DejaVu Sans','axes.spines.top':False,'axes.spines.right':False})
    fig,axs=plt.subplots(1,2,figsize=(7.1,3.65),sharey=True)
    for ax,split,title in zip(axs,['test','x1'],['(a) Same structure   800 rollouts','(b) New structure X1   400 rollouts']):
        for model in colors:
            rs=[r for r in rows if r['split']==split and r['model']==model]
            x=np.array([int(r['k']) for r in rs]);y=np.array([100*float(r['exact']) for r in rs])
            style='--' if model in ['B0_typed','D0_gen','persistence'] else '-'
            ax.plot(x,y,style,marker='o' if model=='A_jev' else '.',lw=1.6,color=colors[model],label=labels[model])
            if rs[0]['ci_lo']:
                ax.fill_between(x,[100*float(r['ci_lo']) for r in rs],[100*float(r['ci_hi']) for r in rs],color=colors[model],alpha=.08)
        ax.set_title(title,fontsize=9.5,pad=11);ax.set_xticks([1,2,3,4]);ax.set_xlim(.9,4.1);ax.set_ylim(0,105)
        ax.set_xlabel('Rollout depth k');ax.grid(axis='y',alpha=.18);ax.set_yticks([0,20,40,60,80,100])
    axs[0].set_ylabel('Exact state match (%)')
    handles,labs=axs[0].get_legend_handles_labels();fig.legend(handles,labs,loc='lower center',ncol=3,frameon=False,fontsize=8.5)
    fig.subplots_adjust(left=.10,right=.98,top=.86,bottom=.28,wspace=.12)
    fig.savefig(QA/'rollout.png',dpi=300,facecolor='white');plt.close(fig)

make_flow();make_rollout()
doc=Document()
sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=Mm(21);sec.right_margin=Mm(21)
sec.header_distance=Mm(8);sec.footer_distance=Mm(9)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    st.element.rPr.rFonts.set(qn('w:ascii'),FONT);st.element.rPr.rFonts.set(qn('w:hAnsi'),FONT)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        st.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    for border in list(st.element.findall('.//'+qn('w:pBdr'))):
        border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=Pt(15);normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
for name,size in [('Title',21),('Subtitle',10.5),('Heading 1',16),('Heading 2',11.5),('Caption',9)]:
    st=doc.styles[name];st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=Pt(26 if name=='Title' else 20 if name=='Heading 1' else 15 if name=='Heading 2' else 12)
doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 2'].font.bold=True
doc.styles['Subtitle'].font.italic=False
doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.core_properties.title='결정 모델을 월드 모델로 활용하는 LLM 에이전트 연구'
doc.core_properties.subject='교수님 검토용 연구 배경 실험 설계 결과 해설'
doc.core_properties.author='송용휘'
doc.core_properties.keywords='JEV, 월드 모델, LLM 에이전트, TextWorld, 결정형, 생성형'
(QA/'fonts.conf').write_text('''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>
<dir>'''+str(QA/'fonts')+'''</dir>
<dir>/Users/song-yonghwi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype</dir>
<dir>/private/tmp/kiis_professor_plotdeps/matplotlib/mpl-data/fonts/ttf</dir>
<cachedir>'''+str(QA/'fontcache')+'''</cachedir>
<alias><family>sans-serif</family><prefer><family>Noto Sans CJK KR</family></prefer></alias>
<alias><family>Cambria Math</family><prefer><family>DejaVu Sans</family></prefer></alias>
</fontconfig>''')

fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run();r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

def p(text,style=None,size=None,bold=False):
    z=doc.add_paragraph(style=style);z.paragraph_format.keep_together=True
    r=z.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return z

def h(text):doc.add_heading(text,level=2)
def page(n,title):
    z=doc.add_heading(f'{n} {title}',level=1)
    if n>1:z.paragraph_format.page_break_before=True

def mathline(text):
    z=doc.add_paragraph();z.alignment=WD_ALIGN_PARAGRAPH.CENTER
    z.paragraph_format.space_before=Pt(3);z.paragraph_format.space_after=Pt(9)
    z.paragraph_format.line_spacing=1.0
    def run(t,upright=False):
        el=OxmlElement('m:r')
        if upright:
            pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);el.append(pr)
        mt=OxmlElement('m:t');mt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');mt.text=t;el.append(mt);return el
    def sub(base,index):
        el=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(base if not isinstance(base,str) else run(base));el.append(e)
        s=OxmlElement('m:sub');s.append(run(index,True));el.append(s);return el
    def hat(base):
        el=OxmlElement('m:acc');pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'̂');pr.append(ch);el.append(pr)
        e=OxmlElement('m:e');e.append(run(base));el.append(e);return el
    om=OxmlElement('m:oMath')
    if text.startswith('T('):
        parts=[run('T(s′ | s, a)'),run(' ;  ',True),sub(hat('s'),'t+1'),run(' = '),sub('f','WM'),run('('),sub('s','t'),run(', '),sub('a','t'),run(')')]
    elif text.startswith('v̂'):
        parts=[sub(hat('v'),'j'),run(' = '),sub(run('arg max',True),'v'),run(' '),sub('q','j'),run('(v | s, a)'),run(' ;  ',True),run('p = '),sub('q','exec'),run('('),run('executes',True),run(' | s, a)')]
    else:
        parts=[run('J(u) = '),run('conj',True),run(' + 0.25 '),run('progress',True),run(' − 0.1 E['),sub('N','invalid'),run('] − 0.01 h + 0.05 '),run('prior',True)]
    for el in parts:om.append(el)
    z._p.append(om)

def tbl(headers,rows,widths,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        q=OxmlElement('w:'+side);q.set(qn('w:val'),'single');q.set(qn('w:sz'),'4');q.set(qn('w:color'),'D9D9D9');borders.append(q)
    pr.append(borders)
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Mm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'75' if side in ['top','bottom'] else '90');el.set(qn('w:type'),'dxa');m.append(el)
            tcpr.append(m)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');tcpr.append(sh)
            for z in c.paragraphs:
                z.paragraph_format.space_after=Pt(0);z.paragraph_format.space_before=Pt(0);z.paragraph_format.line_spacing=Pt(12.5)
                if ci==0 or (len(headers)>2 and ci>0 and len(z.text)<25):z.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in z.runs:r.font.name=FONT;r.font.size=Pt(size);r.bold=ri==0
    blank=p('',size=2);blank.paragraph_format.space_after=Pt(0);blank.paragraph_format.line_spacing=Pt(3)
    return t

def picture(path,caption,width=168):
    z=doc.add_paragraph();z.alignment=WD_ALIGN_PARAGRAPH.CENTER;z.paragraph_format.keep_with_next=True
    z.paragraph_format.line_spacing=1.0
    z.paragraph_format.space_after=Pt(4);shape=z.add_run().add_picture(str(path),width=Mm(width))
    shape._inline.docPr.set('descr',caption)
    z=p(caption,'Caption');z.paragraph_format.space_after=Pt(9)

def ref(num,text,url):
    z=doc.add_paragraph();z.paragraph_format.space_after=Pt(3);z.paragraph_format.line_spacing=Pt(11)
    r=z.add_run(f'[{num}] '+text+' ');r.font.size=Pt(8.5)
    rid=z.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),rid)
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');pr.append(sz)
    color=OxmlElement('w:color');color.set(qn('w:val'),'145F98');pr.append(color);rr.append(pr)
    txt=OxmlElement('w:t');txt.text='원문';rr.append(txt);hyp.append(rr);z._p.append(hyp)

# 1
p('결정 모델을 월드 모델로 활용하는\nLLM 에이전트 연구','Title')
p('연구 배경 실험 설계 결과 해설','Subtitle')
p('송용휘  |  충북대학교 정보통신공학부  |  2026년 10월 2일',size=9.5)
page(1,'연구 개요와 기여')
p('본 연구는 선택지 위의 확률을 반환하는 결정 모델 JEV를 LLM 에이전트의 월드 모델로 활용하고, 그 예측이 실제 행동 선택과 목표 달성에 기여하는지 평가한다. 완전관측 TextWorld 과제에서 과제별 추가 학습 없이 사용한 JEV의 성공률은 93.4%였으며, 실패 기억만 사용하는 기준선보다 13.9%p 높았다. 학습한 LLM은 같은 구조에서 높은 예측 정확도를 보였지만, 새 구조에서는 정확도가 크게 낮아졌다.')
h('연구의 문제 설정')
p('월드 모델의 실용성은 전이 예측의 정확도와 그 예측을 활용한 에이전트 성과를 함께 보아야 한다. 본 연구는 동일한 정책과 계획 절차 아래에서 월드 모델을 교체하여 실제 성과를 비교하고, 별도의 다단계 rollout 평가로 예측 능력과 구조 일반화를 측정한다. 두 평가를 분리하여 정책이나 계획기의 제한을 월드 모델의 오류와 혼동하지 않도록 한다.')
h('방법과 실증적 기여')
p('방법 측면에서는 실행 여부와 상태 변수의 값을 선택지로 질의하고, 답을 조합해 다음 상태를 구성하는 예측 인터페이스를 제시한다. 이를 반복 적용해 후보 행동열의 결과를 예측하고, 목표 달성·진행도·무효 행동 등을 반영한 공통 효용으로 후보를 선택한다.')
p('실증 측면에서는 동일 Qwen3-4B의 결정형·생성형과 학습 전·후를 비교하는 2×2 설계에 JEV를 추가한다. 정책만 사용하는 기준선, 실패 기억 기준선, 실행 여부만 아는 참조 조건, 정확한 전이를 제공하는 oracle을 포함해 성과의 원인을 분석한다.')
h('핵심 결과와 해석')
p('첫째, JEV와 학습한 LLM의 성과 이득은 실패 기억만으로 설명되지 않았다. 둘째, 학습 없는 Qwen의 결정형·생성형 성공률 차이 24.0%p는 생성형 출력의 형식 오류와 그 처리 규칙의 영향으로 해석된다. 셋째, 학습한 LLM의 4단계 예측 정확도는 같은 구조에서 96.8~98.5%였지만 새 구조에서는 66.5~71.5%로 낮아졌다. JEV의 과제 성공률은 oracle과 통계적으로 구분되지 않았으며, 동등성을 입증한 결과로 해석하지 않는다.')
p('연결되는 학회 원고는 「결정 모델을 월드 모델로: LLM 에이전트에서 JEV와 생성형 LLM의 비교」이다. 본 문서는 원고의 연구 배경과 비교 설계, 최종 v3 결과의 의미 및 한계를 상세히 설명한다.',size=9.5)

# 2
page(2,'월드 모델과 모델 기반 의사결정')
h('전이 모델로서의 정의')
p('월드 모델은 에이전트가 환경의 변화 과정을 내부적으로 예측하기 위한 모델이다. 모델 기반 의사결정에서는 현재 상태와 행동에 조건화된 다음 상태를 예측하고, 필요에 따라 보상이나 종료 여부도 모델링한다. 확률적 환경에서는 전이 분포를, 결정적 환경에서는 전이 함수 또는 그 근사를 사용할 수 있다.')
mathline('T(s′ | s, a)            ŝₜ₊₁ = fWM(sₜ, aₜ)')
p('여기서 s는 환경 상태, a는 행동, s′는 실제 다음 상태, ŝ는 예측 상태를 나타낸다. 부분 관측 환경에서는 관측·행동 이력이나 belief state를 모델의 입력으로 사용한다. 본 연구는 현재 상태의 물리적 사실을 제공하는 완전관측·결정적 환경을 대상으로 하며, 숨은 상태 추정의 성능은 평가하지 않는다.')
h('정책 학습과 실행 전 계획')
p('월드 모델은 가상 경험을 생성해 정책을 학습하는 데 활용할 수 있고, 실행 시 후보 행동의 미래 결과를 비교하는 데에도 활용할 수 있다. Ha와 Schmidhuber의 World Models는 환경의 공간적·시간적 표현을 학습하고 모델이 생성한 가상 환경에서 정책을 학습하는 대표 사례다 [1]. 본 연구는 정책의 가중치를 고정한 채, 실행 전 후보 평가에 월드 모델을 사용하는 경우에 초점을 둔다.')
picture(QA/'architecture.png','그림 1  공통 에이전트의 후보 생성 전이 예측 행동 선택 및 재계획 과정')
h('정책과 계획기의 기능 분리')
p('정책은 현재 상태와 목표에서 후보 행동열을 생성한다. 월드 모델은 각 후보의 결과를 예측하고, 계획기는 그 결과의 효용을 비교한다. 본 연구에서는 길이 2까지 예측한 후보 중 하나를 선택하되 첫 행동만 실행하고 실제 관측으로 다시 계획한다. 따라서 예측 오차가 생기더라도 다음 실행 단계에서 실제 상태를 입력받을 수 있다.')
p('정확한 전이 예측을 제공해도 정책이 필요한 후보를 제안하지 않거나 효용이 적절한 후보를 선택하지 못하면 실패할 수 있다. 전이 정확도와 과제 성공률을 별도로 측정하는 이유이며, oracle 역시 동일 후보와 계획기에 묶인 진단용 참조 조건이다.')

# 3
page(3,'LLM 월드 모델의 선행 연구와 쟁점')
p('LLM 기반 월드 모델은 상태 또는 상호작용 이력과 가상의 행동을 입력받아, 다음 환경 상태나 관측을 언어로 예측한다. 기존 모델 기반 접근과 연결되는 지점은 행동에 조건화된 전이 예측이며, 언어 표현을 사용한다는 점 때문에 출력의 일관성과 파싱 가능성이 별도의 성능 요인이 된다.')
h('대표 연구의 접근')
tbl(['연구','주요 접근','이번 연구와의 연결'],[
('World Models [1]','환경 표현과 가상 경험 학습','전이 모델과 정책 활용의 배경'),
('RAP [2]','LLM 예측과 탐색 기반 계획','실행 전 미래 예측을 통한 후보 평가'),
('Wang et al. [3]','텍스트 환경의 상태 전이 예측 평가','에이전트 성과와 별개로 예측 정확도 측정'),
('Xie et al. [4]','실행 조건과 효과 예측 모델의 파인튜닝','환경 데이터 학습의 효과 비교')],[35,66,67],9.3)
p('RAP는 LLM을 월드 모델과 추론 에이전트로 활용하고 MCTS를 통해 추론·행동 경로를 탐색한다 [2]. Wang 등은 ByteSized32 상태 전이 벤치마크에서 GPT-4의 시뮬레이션 신뢰성에 제한이 있음을 보고했다 [3]. Xie 등은 행동의 실행 가능성과 실행 후 상태를 예측하는 두 기능을 각각 파인튜닝하여 LLM의 월드 모델 활용을 연구했다 [4]. 이들 결과를 모든 LLM이나 모든 환경에 대한 결론으로 일반화하지 않는다.')
h('예측 인터페이스와 상태의 일관성')
p('상태를 생성하는 모델은 변경된 사실을 추가하면서 기존 사실을 남길 수 있다. 예를 들어 음식이 탁자와 인벤토리에 동시에 있다고 출력하면, 문자열이 자연스럽더라도 단일 상태로 사용할 수 없다. 반대로 선택지로 상태 변수의 값을 읽으면 변수마다 하나의 값을 얻을 수 있으나, 변수 사이의 공동 제약까지 자동으로 만족하는 것은 아니다.')
h('다단계 예측과 분포 변화')
p('모델 자신의 예측을 다음 입력으로 사용하는 자유 rollout에서는 초기 오류가 후속 입력을 바꾸며 누적될 수 있다. 실제 이전 상태를 입력하는 teacher-forced 예측과 비교하면 이러한 피드백의 영향을 진단할 수 있다. 또한 새로운 이름에 대한 평가와 새로운 환경 구조에 대한 평가는 다른 일반화 문제이므로 구분해야 한다.')
p('본 연구는 결정 모델 JEV의 선택지 확률 출력 [5]을 전이 예측에 연결한다. 동일 Qwen의 결정형·생성형을 함께 평가하여 예측 방식의 차이를 검토하고, JEV와 Qwen 비교는 모델 크기·학습 이력·구현이 함께 다른 시스템 비교로 해석한다. JEV의 내부 학습 목적이나 표현을 성능 차이의 원인으로 식별하지 않는다.')

# 4
page(4,'연구 질문과 비교 설계')
h('연구 질문')
p('RQ1은 월드 모델이 실제 문제 해결에 추가 이득을 주는가이다. 정책만 사용하는 C와 실패 기억을 추가한 C_fm을 구분하여, 반복 실패 억제와 전이 예측의 기여를 나눈다. RQ2는 동일한 Qwen에서 결정형·생성형 예측 방식이 성과와 형식 오류에 어떤 차이를 만드는가이다.')
p('RQ3은 같은 환경 전이로 학습한 LLM이 과제 학습 없이 사용한 JEV와 비교해 어느 수준에 도달하는가이다. RQ4는 다단계 예측과 새 구조에서 정확도가 어떻게 변하는가이다. 연구 질문별로 closed-loop 성과와 독립적인 전이 예측 결과를 연결한다.')
h('비교군의 역할')
tbl(['표기','월드 모델 또는 선택 규칙','질의 방식','과제 학습'],[
('C','정책 1순위 행동','없음','없음'),
('C_fm','정책 순위와 실패 기억','없음','없음'),
('validity','엔진의 실행 가능 여부','정답 참조','없음'),
('Oracle','엔진의 실제 다음 상태','정답 참조','없음'),
('A','JEV jev-1.13.0','결정형','없음'),
('B0','Qwen3-4B','결정형','없음'),
('D0','Qwen3-4B','생성형','없음'),
('B','Qwen3-4B + LoRA','결정형','6,000 전이'),
('D','Qwen3-4B + LoRA','생성형','6,000 전이')],[23,73,37,35],9.5)
h('공정한 비교와 식별 범위')
p('B0와 D0, B와 D는 같은 기반 모델을 사용하므로 예측 프로토콜의 차이를 비교할 수 있다. 다만 결정형은 선택지·확률과 beam을 사용하고 생성형은 단일 상태를 출력하므로, 순수한 질문 문구의 효과로 한정하지 않는다. 후속 분석에서 beam 1 대조와 파싱 오류를 함께 확인한다.')
p('모든 비교군은 동일 정책·후보 생성 설정을 사용한다. 동일 입력과 정책 seed에서는 동일한 후보 생성 규칙이 적용되지만, 행동 선택 이후 상태와 이력이 달라지면 후보도 달라질 수 있다. 월드 모델을 사용하는 계획기 조건에는 동일 효용과 실패 기억을 적용하고, C_fm은 실패하지 않은 후보 중 정책 순위가 높은 행동을 선택한다.')
p('JEV에는 이 과제의 추가 학습을 수행하지 않았다는 의미로 frozen 또는 과제 학습 없음이라고 표현한다. 이는 사전학습 자체가 없다는 뜻이 아니다. A와 B·D의 비교는 범용 모델과 과제 전이로 학습한 모델의 비교이며, 동일 학습 예산 조건의 비교는 아니다.')

# 5
page(5,'제안 방법과 예측 인터페이스')
h('결정형 한 단계 예측')
p('결정형은 상태와 행동 하나에 대해 실행 여부와 동적 상태 변수 6개의 값을 질문한다. 기본 변수는 플레이어 위치, 열쇠·목표 음식·미끼 음식의 소재, 상자 상태, 문 상태다. 물체가 가질 수 있는 위치나 소멸 상태 등을 선택지로 열거하고, 변수별 최빈값으로 실행 후 상태를 구성한다.')
mathline('v̂ⱼ = arg maxᵥ qⱼ(v | s, a)      p = qexec(executes | s, a)')
p('JEV는 API의 선택지 확률을 사용한다. Qwen 결정형은 같은 정보를 질문·선택지로 직렬화하고, 선택지 글자의 다음 토큰 확률을 정규화해 읽는다. 상태 재구성은 변수마다 값을 하나 배정하고 나머지 사실을 유지한다. 전이 규칙을 이용해 예측 오류를 수정하는 별도의 동역학 보정은 하지 않는다.')
tbl(['상태 변수','선택지의 예 또는 역할'],[
('실행 여부','executes / fails'),
('플레이어 위치','현재 방 또는 연결된 다른 방'),
('물체 3개의 소재','인벤토리 상자 탁자 방 바닥 소멸 등'),
('상자 및 문 상태','열림 닫힘 잠김 등 해당 객체의 가능한 상태')],[50,118],9.5)
h('확률 분기와 다단계 rollout')
p('실행 확률 p로 예측 상태에 도달하는 분기와 상태를 유지하는 실패 분기를 가중한다. 분기를 재귀적으로 이어 폭 4의 beam을 유지하고 가중치 0.02 미만의 가지를 제거한다. 변수별 분포와 최빈값 조합은 전체 상태에 대한 보정된 결합분포를 의미하지 않는다. 현재 beam은 주로 실행·실패 경로의 불확실성을 전달한다.')
h('생성형 출력과 파싱 계약')
p('생성형은 result 줄과 다음 상태의 사실 목록을 greedy로 출력한다. 변수의 값이 누락되거나 한 물체의 소재가 둘 이상이면 파싱 실패로 처리한다. 고정 seed의 샘플링으로 한 번 재생성하고, 여전히 실패하면 closed-loop에서는 실행 확률 0의 상태 유지 예측으로 처리한다. 출력 길이는 최대 320 토큰이며, dev에서 고정한 형식 예시 하나를 사용한다.')
p('예를 들어 take apple from table 이후 on(apple, table)과 in(apple, inventory)를 함께 생성하면 소재를 하나로 결정할 수 없다. 결정형은 apple의 소재 질문에서 하나를 읽어 상태를 구성한다. 이러한 형식적 차이가 행동 선택에 미치는 영향은 9절에서 분석한다.')
h('공통 후보 평가')
mathline('J(u) = conj + 0.25 progress − 0.1 E[Ninvalid] − 0.01 h + 0.05 prior')
p('conj는 목표 원자의 동시 충족, progress는 충족 비율, h는 후보 길이, prior는 정책 순위에서 얻는 선호다. 계수는 공통으로 고정했다. 현재 상태에서 실패한 첫 행동의 후보를 제외한 뒤 효용이 높은 후보의 첫 행동을 실행한다.',size=9.8)

# 6
page(6,'실험 프로토콜과 재현성')
h('환경과 데이터')
p('TextWorld 1.7.0 symbolic JSON backend [6]에서 방 두 개, 열쇠가 든 상자, 잠긴 문, 탁자 위의 목표 음식과 미끼 음식을 구성했다. 목표는 목표 음식을 소지하고 문을 연 상태에서 먼 방에 도달하는 것이다. 최단 해결 경로는 상자 열기, 열쇠 집기, 문 잠금 해제, 목표 음식 집기, 문 열기, 이동의 6개 행동이다.')
tbl(['항목','최종 평가 설정'],[
('데이터 분할','dev 12 world train 150 world val 20 world test 24 world'),
('학습 전이','B·D 동일 6,000개 엔진 정답 전이'),
('정책','Qwen3-4B bf16 thinking 끔 T=0.7 top_p=0.9'),
('후보 생성','길이 2 후보를 샘플당 최대 8개 두 샘플 병합'),
('계획 및 종료','H=2 30 step 이내 목표 충족 시 성공'),
('평가 규모','24 world × 4 시작 상태 × 3 정책 seed = arm당 288'),
('LoRA 학습','r=16 α=32 dropout 0.05 lr=10⁻⁴ 2 epoch'),
('Checkpoint 선택','val 전이 200개의 정확도로 선택 학습 seed 0'),
('예측 평가','같은 구조 800 새 구조 X1 400 rollout 길이 4'),
('신뢰구간','world 단위 paired bootstrap 4,000회 95% CI')],[39,129],9.1)
p('B는 전이당 질문 7개의 선택지 글자를 학습하고, D는 실행 여부와 다음 상태 사실 목록을 학습한다. 정답은 엔진에서 얻으며 JEV 출력은 학습이나 checkpoint 선택에 사용하지 않는다. 정책 생성 시에는 LoRA를 끄고 월드 모델 추론 시에만 켠다. v3는 val 점검을 통과한 v1 학습 어댑터를 재사용했다.')
h('평가 지표와 통계 단위')
p('과제 성공률은 30 step 안에 목표 원자 3개를 만족한 episode 비율이다. 무효율은 실행되지 않은 행동 수를 전체 실행 step으로 나눈다. 상태 완전일치는 각 깊이에서 평가 대상 변수가 모두 정답과 같은 경우로 정의한다. 같은 구조는 6개, X1은 9개 변수를 평가한다. episode는 같은 world·시작 상태·정책 seed로 짝짓고, CI는 288 episode가 아니라 24 world를 재표집하여 계산한다.')
h('보정 이력과 결과 추적')
p('v1은 후보 부족으로 oracle과 좋은 월드 모델의 성과가 함께 제한되었다. v2 보정은 test에 충분히 이전되지 않아 pilot으로 보존했다. v3는 일반 행동 규칙 안내, 후보 두 샘플 병합, 반복 실패 시 재제안 규칙을 공통 적용하고 새 test 어휘를 사용했다. 이전 v2 test는 보정용으로 전환했으며, 학습·보정과 최종 test 이름을 분리했다.')
p('최종 조건 해시는 448b4fac3efe이다. 로그에는 조건·arm·코드 revision·world 지문을 기록하고 조건이 다른 결과를 합치지 않는다. 본 문서의 수치는 results/slots.json에서 가져왔으며, 원시 로그 재계산 검증을 통과했다.',size=9.5)

# 7
page(7,'에이전트 성과와 Closed loop 평가')
p('최종 v3 평가는 9개 arm 각각 288 episode로 총 2,592 episode를 실행했다. 표 1은 실패 기억의 효과와 월드 모델의 추가 효과를 함께 보여 준다. 대괄호는 world 단위 95% bootstrap 신뢰구간이다.')
rows=[]
for label,key,slot in [('C','C','R1'),('C_fm','C_fm','R15'),('validity','validity',None),('Oracle','oracle',None),('A  JEV','A_jev','R2'),('B0  결정형','B0_typed','R3'),('D0  생성형','D0_gen','R4'),('B  결정형 학습','B_typed','R5'),('D  생성형 학습','D_gen','R6')]:
    v=(S['R15']['C_fm'] if slot=='R15' else S[slot]) if slot else S['R7'][key]
    rows.append((label,f"{v['text']} {v['ci_text']}",S['R9']['values'][key]))
tbl(['비교군','성공률 % 및 95% CI','무효율 %'],rows,[49,84,35],9.7)
p('표 1  과제 성공률과 실제 실행 행동의 무효율','Caption')
h('월드 모델의 추가 효과')
p('C의 성공률은 17.0%, C_fm은 79.5%로 반복 실패 억제의 영향이 컸다. 이를 통제한 뒤에도 A−C_fm은 +13.9%p [2.1, 24.3], B−C_fm과 D−C_fm은 각각 +13.2%p [2.8, 22.9]였다. 따라서 JEV와 학습한 월드 모델의 이득은 실패 기억만으로 설명되지 않는다. 실행 여부만 제공한 validity 85.8%와 oracle 92.7%의 차이도 +6.9%p [1.0, 13.9]로, 실제 결과 예측이 더하는 효과가 있었다.')
h('예측 방식과 학습의 영향')
p('학습 없는 Qwen에서 B0−D0는 +24.0%p [11.5, 35.8]였다. D0는 실패 기억 기준선보다 20.5%p 낮았고, 형식 오류로 유용한 행동을 회피하는 사례가 관찰되었다. 학습 후 B와 D는 모두 92.7%였으며 288 episode 전부 oracle과 성공 여부가 일치했다. 방식 차이는 학습 후 사라졌지만, 이는 모든 환경에서 두 방식이 같은 성능이라는 뜻은 아니다.')
h('Oracle과 통계적 해석')
p('A의 93.4%와 oracle의 92.7% 차이는 +0.7%p [−1.7, 3.1]로 구분되지 않았다. A−B0도 +10.4%p [−1.7, 21.5]로 구분되지 않았다. CI가 0을 포함하는 비교를 우월성이나 동등성의 증거로 표현하지 않는다. 정확한 예측을 제공한 oracle도 21개 episode에서 실패했으며, 후보 부족 6개와 계획기 관련 실패 15개로 분류되었다. 좋은 모델 간 구분은 다음 절의 전이 예측 평가에서 확인한다.')

# 8
page(8,'다단계 예측과 구조 일반화')
p('자유 rollout은 예측한 상태를 다음 입력으로 이어 사용한다. 평가 행동열은 정책 계획과 무작위 행동열 두 종류이며, 같은 구조에서 800개, 새 구조 X1에서 400개를 사용했다. 생성형 파싱 실패는 해당 step부터 이후 깊이까지 오답으로 처리한다.')
picture(QA/'rollout.png','그림 2  자유 rollout의 깊이별 상태 완전일치율 음영은 95% CI')
h('학습과 같은 구조에서의 예측')
p('4단계 완전일치율은 A 77.9%, B0 9.1%, D0 6.2%, B 98.5%, D 96.8%였다. 학습 없는 Qwen은 상태가 바뀌지 않는다고 예측하는 persistence 15.6%보다 낮았다. 학습한 B·D는 JEV보다 정확했지만, closed-loop 성공률은 oracle과 일치하여 예측 정확도의 추가 이득이 최종 성과 차이로 드러나지 않았다.')
h('이름 분할과 구조 일반화의 구분')
p('최종 test는 이름이 새롭지만 과제 구조는 학습과 같다. 이름을 객체 id로 치환하면 평가 전이 2,869개 중 1,995개인 69.5%가 실제 학습 전이 6,000개에 존재한다. 따라서 이 평가의 높은 정확도는 새 이름에 대한 전이 능력과 학습한 구조의 재현 능력을 함께 반영하며, 새로운 과제 구조에 대한 일반화 근거로 사용하지 않는다.')
p('X1은 곁방, 미끼 열쇠, 음식이 든 둘째 용기를 추가한 구조로, 4단계 정확도는 A 85.5%, B 71.5%, D 66.5%였다. world별 k=1~4 평균 정확도의 test−X1 변화량을 JEV와 비교한 차이는 B +25.6%p [18.8, 33.9], D +26.7%p [21.7, 31.6]였다. 이는 4단계 하나의 하락폭이 아니라 깊이 평균 변화량의 짝비교다.')
p('JEV의 X1 정확도가 더 높은 현상은 무효 행동과 상태 유지 step이 더 많은 행동열 구성의 영향을 받는다. 두 조건의 절대 정확도를 단순히 난이도 순위로 읽지 않는다. X1은 9개 변수의 예측을 평가한 결과이며, 새 구조의 closed-loop 성공률은 측정하지 않았다.')

# 9
page(9,'오류 분석과 결과 해석')
h('형식 오류와 실제 행동 선택')
p('D0는 음식 집기 이후 새 위치를 쓰면서 옛 위치를 남기는 오류가 빈번했다. 재생성 후에도 파싱되지 않으면 실행 안 됨으로 처리하므로, 이 후보는 무효 행동 벌점을 받는다. 음식 없이 먼 방에 진입한 episode는 D0 115/288, C_fm 61/288이었고, 진입 이후 성공은 각각 0/115와 3/61이었다. 이는 환경의 비가역 실패로 확인된 현상이라기보다, 음식 확보 행동을 다시 제안하지 못하는 정책의 회복 한계와 연결된다.')
p('한 단계 teacher-forced 예측 기준 파싱 실패율은 D0 26.0% (833/3,200), D 0.1% (3/3,200)였고 X1에서는 46.9%와 2.1%였다. closed-loop D0 로그는 예측 수 대신 요청 수를 기록하여 실패율을 26.1~41.4% 범위로만 산출할 수 있다. 26.0%를 closed-loop의 정확한 실패율로 사용하지 않는다.')
p('D0가 파싱된 동일 행에서 B0와 D0의 정확도는 각각 37.1%와 46.2%였다. 파싱 가능한 행의 선택 편향이 있으므로 생성형의 내용 우월성을 주장할 수 없지만, 결정형이 전이 내용을 더 잘 예측한다는 근거도 없다. 관찰된 closed-loop 이점은 형식 견고성의 영향으로 해석하며, 파싱 실패 처리 규칙에 대한 민감도 재실험은 수행하지 않았다.')
h('오류 누적과 Beam 대조')
p('실제 이전 상태에서 한 단계씩 예측하는 teacher-forced와 자유 rollout의 차이를 k=2~4에서 평균한 누적 격차 G를 계산했다. 같은 구조의 G 차이 D0−B0는 −1.7%p [−4.5, 1.0], D−B는 +0.8%p [−0.1, 1.9]로 구분되지 않았다. X1과 파싱 실패를 제외한 분석에서도 방식 차이를 확인하지 못했다. 따라서 결정형이 내용 오류의 누적을 억제한다는 주장은 현재 결과로 뒷받침되지 않는다.')
p('Qwen 결정형의 beam 1과 beam 4 대조는 1위 상태 예측에 거의 차이를 만들지 않았다. 이 결과는 해당 예측 평가의 관찰이며, 모든 계획 상황에서 beam의 효과가 없음을 의미하지 않는다. JEV는 같은 rollout 반복에서도 일부 응답이 달라 외부 API의 비결정성을 재현성 한계로 기록한다.')
h('계산 비용')
tbl(['모델','월드 모델 추론 비용 또는 시간','학습 GPU h'],[
('A  JEV','episode당 129 요청 약 $0.010','과제 학습 없음'),
('B0 / D0','episode당 49 / 115 GPU 초','없음'),
('B / D','episode당 42 / 73 GPU 초','13.9 / 4.5')],[33,92,43],9.1)
p('표 2  Closed loop 월드 모델 비용 학습 시간은 별도 집계','Caption')
p('API 비용과 GPU 시간을 같은 통화 비용으로 환산하지 않는다. 정책 비용을 제외한 월드 모델 사용량이며, 학습 비용과 외부 서비스 지연·가격 정책까지 포함한 총비용의 우열은 이 표로 판단할 수 없다.',size=9.5)

# 10
page(10,'결론 타당성 한계 및 후속 연구')
h('연구 질문에 대한 답')
p('JEV와 학습한 LLM 월드 모델은 실패 기억을 넘어 과제 성공률을 높였다. 동일 Qwen의 학습 전 방식 차이는 출력 형식과 실패 처리의 영향이 컸고, 학습 후에는 closed-loop 차이가 사라졌다. 학습한 LLM은 같은 구조에서 가장 정확했으나 새 구조에서는 JEV보다 낮았다. 다단계 예측을 통해 모델 간 차이를 확인했지만 결정형의 내용 오류 누적 억제는 확인하지 못했다.')
h('타당성 한계')
p('단일 Qwen 모델과 학습 seed, 한 가지 기본 과제 구조, 연구자가 설계한 상태 변수 스키마에서 얻은 결과다. JEV와 Qwen의 학습 이력·구조 차이를 통제하지 못했고, 외부 API의 비결정성과 버전 의존성이 남는다. 형식 실패 처리 규칙의 민감도와 새 구조의 실제 에이전트 성과도 추가 검증이 필요하다.')
h('후속 연구와 검토 사항')
p('후속 계획은 ScienceWorld를 주 환경, ALFWorld를 추가 환경으로 삼아 관측 이력 기반 예측과 실제 문제 해결을 평가하는 것이다. 동일 LLM의 자유 생성·스키마 제약 생성·결정형을 비교해 형식 효과를 분리하고, 공통 입력·정책·계산 예산 아래에서 JEV 및 공개 월드 모델을 평가한다. 현재는 구현 전 계획이다. 교수님과는 기여의 범위, 추가 통제 실험의 우선순위, 더 큰 논문으로 확장할 핵심 질문을 논의하고자 한다.')
h('참고문헌')
ref(1,'D. Ha and J. Schmidhuber. World Models. arXiv:1803.10122, 2018.','https://arxiv.org/abs/1803.10122')
ref(2,'S. Hao et al. Reasoning with Language Model is Planning with World Model. EMNLP, pp. 8154–8173, 2023.','https://aclanthology.org/2023.emnlp-main.507/')
ref(3,'R. Wang et al. Can Language Models Serve as Text-Based World Simulators? ACL Short Papers, pp. 1–17, 2024.','https://aclanthology.org/2024.acl-short.1/')
ref(4,'K. Xie et al. Making Large Language Models into World Models with Precondition and Effect Knowledge. COLING, pp. 7532–7545, 2025.','https://aclanthology.org/2025.coling-main.503/')
ref(5,'TypeSafe AI. Introducing System One Models & Jev. 공식 기술 소개, 열람 2026년 10월 2일.','https://typesafe.ai/blog/introducing-system-one-models-and-jev')
ref(6,'M.-A. Côté et al. TextWorld: A Learning Environment for Text-based Games. arXiv:1806.11532, 2018.','https://arxiv.org/abs/1806.11532')
ref(7,'A. Yang et al. Qwen3 Technical Report. arXiv:2505.09388, 2025.','https://arxiv.org/abs/2505.09388')
ref(8,'E. J. Hu et al. LoRA: Low-Rank Adaptation of Large Language Models. ICLR, 2022.','https://arxiv.org/abs/2106.09685')
p('실험 근거  results/slots.json 및 fig2.csv  |  조건 448b4fac3efe\n실행 로그  artifacts/kiis_k4v3  kiis_k5v3  kiis_x1v3  |  학습 기록  artifacts/kiis_k3',size=8.2)

# Add citations for model and parameter-efficient tuning where actually used.
for z in doc.paragraphs:
    if z.text.startswith('B는 전이당 질문 7개'):
        z.add_run(' 기반 모델과 LoRA의 일반적 배경은 [7, 8]을 참조한다.').font.size=Pt(10.5)

out=BASE/'교수님_검토용_연구설명서.docx'
doc.save(out)
print(out)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
