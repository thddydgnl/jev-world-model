# -*- coding: utf-8 -*-
"""교수님 검토용 연구 설명서 v2 — 「LLM 에이전트의 월드 모델: 결정 모델 JEV와 범용 LLM의 비교」

원본(build_report.py, 2026-10-02)을 바탕으로 구조·서식을 다시 잡고, 배경(강화학습·LLM 에이전트의 월드 모델)과
방법 설명을 보강한 판. A4, 여백 12.7 mm(원본 Word 편집본과 같음), 글꼴 Apple SD Gothic Neo.
수치는 results/slots.json (revision 3ad24a1, 조건 448b4fac3efe)의 v3 값이며, 아래 CHECK에서 대조한다.

사용:  .venv/bin/python kiis2026f/professor_report/build_report_v2.py
출력:  kiis2026f/professor_report/LLM_에이전트의_월드모델_v2.docx
"""
from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
QA = BASE / '.qa'
QA.mkdir(exist_ok=True)
OUT = BASE / 'LLM_에이전트의_월드모델_v2.docx'

FONT = 'Apple SD Gothic Neo'
FONT_FILE = '/System/Library/Fonts/AppleSDGothicNeo.ttc'   # index 0 Regular, 4 SemiBold, 6 Bold
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)

# ---------------------------------------------------------------- 수치 대조
S = json.loads((ROOT / 'results/slots.json').read_text())['slots']
CHECK = {
    'R1': '17.0', 'R2': '93.4', 'R3': '83.0', 'R4': '59.0', 'R5': '92.7', 'R6': '92.7', 'R20': '69.5',
}
for k, v in CHECK.items():
    assert S[k]['text'] == v, (k, S[k]['text'], v)
assert S['R15']['C_fm']['text'] == '79.5'
assert S['R7']['validity']['text'] == '85.8'
assert S['R12']['teacher_forced']['test']['D0_gen']['text'] == '26.0'
assert S['R11']['curves']['A_jev']['4']['text'] == '77.9'

# ---------------------------------------------------------------- 그림
def _font(size, index=0):
    return ImageFont.truetype(FONT_FILE, size, index=index)


def make_world():
    """그림 1: 과제 구성 (두 방, 열쇠가 든 상자, 탁자 위 음식 둘, 잠긴 문)."""
    W, H = 1900, 560
    im = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(im)
    big, mid, small = _font(42, 4), _font(35), _font(31)
    # 방 A
    d.rounded_rectangle((40, 40, 900, 420), radius=14, fill='#f7f8fa', outline='#777777', width=3)
    d.text((70, 58), '시작 방', font=big, fill='#111111')
    # 상자
    d.rounded_rectangle((110, 140, 450, 360), radius=12, fill='#ffffff', outline='#888888', width=2)
    d.text((280, 180), '상자 (닫힘)', font=mid, fill='#111111', anchor='mm')
    d.rounded_rectangle((170, 230, 390, 320), radius=10, fill='#e8edf2', outline='#999999', width=2)
    d.text((280, 275), '열쇠', font=mid, fill='#111111', anchor='mm')
    # 탁자
    d.rounded_rectangle((500, 140, 850, 360), radius=12, fill='#ffffff', outline='#888888', width=2)
    d.text((675, 180), '탁자', font=mid, fill='#111111', anchor='mm')
    d.rounded_rectangle((540, 225, 810, 275), radius=10, fill='#e8edf2', outline='#999999', width=2)
    d.text((675, 250), '목표 음식', font=small, fill='#111111', anchor='mm')
    d.rounded_rectangle((540, 290, 810, 340), radius=10, fill='#f1f3f5', outline='#999999', width=2)
    d.text((675, 315), '미끼 음식', font=small, fill='#111111', anchor='mm')
    # 문
    d.rectangle((900, 170, 1000, 290), fill='#d9dee5', outline='#777777', width=3)
    d.text((950, 212), '문', font=mid, fill='#111111', anchor='mm')
    d.text((950, 255), '잠김', font=small, fill='#333333', anchor='mm')
    # 방 B
    d.rounded_rectangle((1000, 40, 1860, 420), radius=14, fill='#f7f8fa', outline='#777777', width=3)
    d.text((1030, 58), '건너편 방', font=big, fill='#111111')
    d.text((1430, 200), '목표 상태', font=mid, fill='#111111', anchor='mm')
    d.text((1430, 250), '목표 음식을 들고', font=small, fill='#333333', anchor='mm')
    d.text((1430, 292), '문이 열린 채', font=small, fill='#333333', anchor='mm')
    d.text((1430, 334), '이 방에 있기', font=small, fill='#333333', anchor='mm')
    # 아래 설명
    d.text((950, 490), '최단 해법 6행동:  상자 열기 → 열쇠 집기 → 문 잠금 해제 → 목표 음식 집기 → 문 열기 → 건너편 방으로 이동',
           font=small, fill='#333333', anchor='mm')
    im.save(QA / 'world.png')


def make_flow():
    """그림 2: 에이전트의 한 step (원본 그림 1과 같은 그림)."""
    im = Image.new('RGB', (1900, 390), 'white')
    d = ImageDraw.Draw(im)
    f, small = _font(38), _font(29)
    labels = [('현재 상태', '완전관측'), ('정책', '후보 행동열'), ('월드 모델', '전이 예측'),
              ('계획기', '효용 비교'), ('실제 환경', '첫 행동 실행')]
    for i, (a, b) in enumerate(labels):
        x = 20 + i * 380
        d.rounded_rectangle((x, 45, x + 320, 205), radius=12, fill='#f1f3f5', outline='#777777', width=2)
        d.text((x + 160, 99), a, font=f, fill='black', anchor='mm')
        d.text((x + 160, 157), b, font=small, fill='#333333', anchor='mm')
        if i < 4:
            d.line((x + 330, 125, x + 365, 125), fill='#444444', width=4)
            d.polygon([(x + 365, 125), (x + 350, 116), (x + 350, 134)], fill='#444444')
    d.line([(1700, 220), (1700, 280), (180, 280), (180, 220)], fill='#555555', width=3)
    d.polygon([(180, 220), (171, 237), (189, 237)], fill='#555555')
    d.text((950, 326), '새 관측으로 갱신하고 다음 step에서 재계획', font=small, fill='#333333', anchor='mm')
    im.save(QA / 'architecture.png')


make_world()
make_flow()
ROLLOUT = QA / 'rollout.png'          # build_report.py가 results/fig2.csv로 그린 그림 (그대로 사용)
EXAMPLE = ROOT / 'paper/fig1.png'     # 학회 원고의 그림 1 (결정형·생성형 한 step 예측 예)
assert ROLLOUT.exists() and EXAMPLE.exists()

# ---------------------------------------------------------------- 문서 기본 설정
doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.top_margin = Mm(12.7); sec.bottom_margin = Mm(14); sec.left_margin = Mm(12.7); sec.right_margin = Mm(12.7)
sec.header_distance = Mm(7); sec.footer_distance = Mm(7)

styles_el = doc.styles.element
rpr_default = styles_el.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
rfonts = rpr_default.find(qn('w:rFonts'))
for a in list(rfonts.attrib):
    del rfonts.attrib[a]
for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
    rfonts.set(qn(a), FONT)
lang = rpr_default.find(qn('w:lang'))
if lang is None:
    lang = OxmlElement('w:lang'); rpr_default.append(lang)
lang.set(qn('w:val'), 'en-US'); lang.set(qn('w:eastAsia'), 'ko-KR')


def set_font(style, size=None, bold=None, color=INK):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if color is not None: style.font.color.rgb = color
    for border in list(style.element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)


def para_flags(ppr_parent):
    """한글 문서 관례: 라틴 단어 중간에서 줄을 바꾸지 않고, 한글·영문 사이 자동 간격을 넣지 않는다."""
    ppr = ppr_parent.get_or_add_pPr()
    for tag in ('w:wordWrap', 'w:autoSpaceDE', 'w:autoSpaceDN'):
        el = OxmlElement(tag); el.set(qn('w:val'), '0'); ppr.append(el)


normal = doc.styles['Normal']
set_font(normal, 10, False)
normal.paragraph_format.line_spacing = Pt(15.2)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control = True
para_flags(normal.element)

for name, size, before, after, line in [('Heading 1', 13.5, 15, 5, 19), ('Heading 2', 11, 9, 3, 16)]:
    st = doc.styles[name]
    set_font(st, size, True)
    pf = st.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = Pt(line)
    pf.keep_with_next = True; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para_flags(st.element)
h1 = doc.styles['Heading 1']
pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '4'); bot.set(qn('w:space'), '2'); bot.set(qn('w:color'), 'BFBFBF')
pbdr.append(bot); h1.element.get_or_add_pPr().append(pbdr)

cap = doc.styles['Caption']
set_font(cap, 9, False, MUTED)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(2); cap.paragraph_format.space_after = Pt(9)
cap.paragraph_format.line_spacing = Pt(12)
cap.font.italic = False

for name in ('List Bullet', 'List Number'):
    st = doc.styles[name]
    set_font(st, 10, False)
    st.paragraph_format.line_spacing = Pt(15.2)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para_flags(st.element)

doc.core_properties.title = 'LLM 에이전트의 월드 모델: 결정 모델 JEV와 범용 LLM의 비교'
doc.core_properties.subject = '교수님 검토용 연구 설명서 (배경, 비교 설계, 결과 해설)'
doc.core_properties.author = '송용휘'
doc.core_properties.keywords = 'JEV, 월드 모델, LLM 에이전트, TextWorld, 결정형, 생성형'

# 바닥글 쪽 번호
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(); r.font.size = Pt(9); r.font.color.rgb = MUTED
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r._r.addnext(fld)


# ---------------------------------------------------------------- 도우미
def p(text, size=None, bold=False, align=None, after=None, before=None, color=None, keep=False, italic=False):
    z = doc.add_paragraph()
    r = z.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    if align is not None: z.alignment = align
    if after is not None: z.paragraph_format.space_after = Pt(after)
    if before is not None: z.paragraph_format.space_before = Pt(before)
    if keep: z.paragraph_format.keep_with_next = True
    return z


def rich(parts, after=None):
    """parts: [(text, bold), ...]"""
    z = doc.add_paragraph()
    for text, bold in parts:
        r = z.add_run(text); r.bold = bold
    if after is not None: z.paragraph_format.space_after = Pt(after)
    return z


def h1_(text): return doc.add_heading(text, level=1)
def h2_(text): return doc.add_heading(text, level=2)


def bullets(items, numbered=False):
    style = 'List Number' if numbered else 'List Bullet'
    for it in items:
        z = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = z.add_run(it[0]); r.bold = True
            z.add_run(it[1])
        else:
            z.add_run(it)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def tbl(headers, rows, widths, size=9, align_first_center=True, center_cols=()):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        q = OxmlElement('w:' + side); q.set(qn('w:val'), 'single'); q.set(qn('w:sz'), '4'); q.set(qn('w:color'), 'D0D5DB')
        borders.append(q)
    pr.append(borders)
    for col, w in zip(t.columns, widths):
        col.width = Mm(w)
    for i, txt in enumerate(headers):
        t.rows[0].cells[i].text = txt
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        cells = t.add_row().cells
        for c, txt in zip(cells, row):
            c.text = str(txt)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, c in enumerate(row.cells):
            c.width = Mm(widths[ci]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '60' if side in ['top', 'bottom'] else '85'); el.set(qn('w:type'), 'dxa'); m.append(el)
            tcpr.append(m)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'E8EDF2'); tcpr.append(sh)
            for z in c.paragraphs:
                z.paragraph_format.space_after = Pt(0); z.paragraph_format.space_before = Pt(0)
                z.paragraph_format.line_spacing = Pt(12.5)
                para_flags(z._p)
                if ri == 0 or (align_first_center and ci == 0) or ci in center_cols:
                    z.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    z.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in z.runs:
                    r.font.size = Pt(size); r.bold = (ri == 0)
                if ri < len(t.rows) - 1 and len(t.rows) <= 12:
                    z.paragraph_format.keep_with_next = True
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(5)
    spacer.add_run('').font.size = Pt(2)
    return t


def picture(path, caption, width=170):
    z = doc.add_paragraph(); z.alignment = WD_ALIGN_PARAGRAPH.CENTER; z.paragraph_format.keep_with_next = True
    z.paragraph_format.line_spacing = 1.0; z.paragraph_format.space_after = Pt(3); z.paragraph_format.space_before = Pt(4)
    shape = z.add_run().add_picture(str(path), width=Mm(width))
    shape._inline.docPr.set('descr', caption)
    doc.add_paragraph(caption, style='Caption')


def mathline(which):
    """OMML 수식 (원본과 같은 식 3개)."""
    z = doc.add_paragraph(); z.alignment = WD_ALIGN_PARAGRAPH.CENTER
    z.paragraph_format.space_before = Pt(2); z.paragraph_format.space_after = Pt(7); z.paragraph_format.line_spacing = Pt(16)

    def run(t, upright=False):
        el = OxmlElement('m:r')
        if upright:
            pr = OxmlElement('m:rPr'); sty = OxmlElement('m:sty'); sty.set(qn('m:val'), 'p'); pr.append(sty); el.append(pr)
        mt = OxmlElement('m:t'); mt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); mt.text = t; el.append(mt)
        return el

    def sub(base, index):
        el = OxmlElement('m:sSub'); e = OxmlElement('m:e'); e.append(base if not isinstance(base, str) else run(base)); el.append(e)
        s = OxmlElement('m:sub'); s.append(run(index, True)); el.append(s); return el

    def hat(base):
        el = OxmlElement('m:acc'); pr = OxmlElement('m:accPr'); ch = OxmlElement('m:chr'); ch.set(qn('m:val'), '̂'); pr.append(ch); el.append(pr)
        e = OxmlElement('m:e'); e.append(run(base)); el.append(e); return el

    om = OxmlElement('m:oMath')
    if which == 1:
        parts = [run('T(s′ | s, a)'), run(' ;    ', True), sub(hat('s'), 't+1'), run(' = '), sub('f', 'WM'), run('('), sub('s', 't'), run(', '), sub('a', 't'), run(')')]
    elif which == 2:
        parts = [sub(hat('v'), 'j'), run(' = '), sub(run('arg max', True), 'v'), run(' '), sub('q', 'j'), run('(v | s, a)'), run(' ;    ', True),
                 run('p = '), sub('q', 'exec'), run('('), run('executes', True), run(' | s, a)')]
    else:
        parts = [run('J(u) = '), run('conj', True), run(' + 0.25 '), run('progress', True), run(' − 0.1 E['), sub('N', 'invalid'), run('] − 0.01 h + 0.05 '), run('prior', True)]
    for el in parts:
        om.append(el)
    z._p.append(om)


# ================================================================ 본문
# ---- 제목
z = p('LLM 에이전트의 월드 모델:', size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, before=14)
z.paragraph_format.line_spacing = Pt(28)
z = p('결정 모델 JEV와 범용 LLM의 비교', size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
z.paragraph_format.line_spacing = Pt(28)
p('연구 배경 · 비교 설계 · 결과 해설', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, color=MUTED)
p('송용휘  |  충북대학교 정보통신공학부  |  2026년 10월 7일', size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, color=MUTED)

# ---- 1
h1_('1. 연구 개요')
p('LLM 에이전트는 행동의 결과를 미리 알지 못한 채 행동을 고르기 때문에, 실행되지 않는 명령이나 목표와 무관한 명령을 되풀이하는 일이 잦다. '
  '행동의 결과를 미리 예측하는 월드 모델을 붙이면 이 문제를 줄일 수 있는데, 지금까지의 LLM 월드 모델은 대부분 다음 상태를 문장으로 생성하는 방식이었다. '
  '본 연구는 선택지 위의 확률로 답하는 결정 모델 JEV를 월드 모델로 쓰는 방법을 제안하고, 같은 에이전트 안에서 월드 모델만 바꿔 가며 '
  '범용 LLM(Qwen3-4B)의 학습 전·후, 그리고 결정형·생성형 질의 방식과 비교했다. 환경은 완전관측 TextWorld 과제이며, 모든 비교군은 같은 정책, '
  '같은 후보 행동열, 같은 계획기를 공유한다.')
p('핵심 결과는 다음 세 가지다.', after=3)
bullets([
    '과제 학습 없이 쓴 JEV의 성공률은 93.4%로, 실패한 행동만 기억하는 기준선(79.5%)보다 13.9%p 높았고, 환경의 정답 전이를 넣은 oracle(92.7%)과 같은 수준이었다.',
    '학습 없는 LLM은 질문을 선택지로 바꾸기만 해도 성공률이 59.0%에서 83.0%로 올랐다. 차이는 생성형 출력의 26.0%가 파싱되지 않는 형식 오류에서 비롯됐다.',
    '환경 전이 6,000개로 파인튜닝한 LLM은 학습한 구조에서는 가장 정확했지만(4 step 예측 98.5%), 학습에 없던 구조에서는 JEV보다 크게 떨어졌다(71.5% 대 85.5%).',
])
p('이 문서는 KIIS 2026 추계학술대회 투고 원고 「결정 모델을 월드 모델로: LLM 에이전트에서 JEV와 생성형 LLM의 비교」(A4 2쪽)의 배경과 설계, 결과를 '
  '풀어 쓴 것이다. 2절은 월드 모델의 배경, 3절은 연구 질문과 비교 설계, 4·5절은 과제 환경과 방법, 6·7절은 실험 설정과 결과, 8절은 해석과 한계, 향후 계획이다.')

# ---- 2
h1_('2. 배경: 에이전트와 월드 모델')
h2_('2.1 월드 모델이란')
p('월드 모델은 에이전트가 환경을 직접 건드리지 않고도 "이 상태에서 이 행동을 하면 무엇이 달라지는가"를 내부적으로 예측하는 모델이다. '
  '가장 단순한 형태는 현재 상태 s와 행동 a를 받아 다음 상태 s′를 예측하는 전이 모델이며, 필요하면 보상이나 종료 여부를 함께 예측한다. '
  '확률적 환경에서는 전이 분포를, 결정적 환경에서는 전이 함수의 근사를 학습하거나 가져다 쓴다.')
mathline(1)
p('여기서 sₜ는 시점 t의 환경 상태, aₜ는 행동, ŝₜ₊₁은 예측한 다음 상태다. 부분 관측 환경이라면 상태 대신 관측·행동 이력이나 belief state가 입력이 된다. '
  '본 연구는 환경이 현재의 물리적 사실을 모두 알려 주는 완전관측·결정적 환경을 대상으로 하므로, 숨은 상태를 추정하는 문제는 다루지 않고 '
  '"행동 뒤에 상태가 어떻게 바뀌는가"의 예측에만 집중한다.')
p('월드 모델의 쓰임새는 크게 둘이다. 하나는 모델이 만들어 내는 가상 경험으로 정책을 학습하는 것이고, 다른 하나는 실행 시점에 후보 행동의 결과를 '
  '미리 내다보고 고르는 것이다. 본 연구는 후자, 즉 정책은 고정한 채 실행 전 후보 평가에 월드 모델을 쓰는 경우를 다룬다.')

h2_('2.2 강화학습 에이전트에서의 월드 모델')
p('월드 모델은 강화학습에서 오래된 개념이다. Sutton의 Dyna [1]는 실제 경험으로 전이 모델을 학습하고, 그 모델이 만들어 낸 가상 경험으로 가치 함수를 '
  '추가로 갱신해 표본 효율을 높였다. 경험을 모델에 담아 두었다가 다시 꺼내 쓴다는 이 발상이 이후 모델 기반 강화학습의 출발점이 됐다. '
  'Ha와 Schmidhuber의 World Models [2]는 화면 관측을 변분 오토인코더로 압축하고 순환 신경망으로 다음 잠재 상태를 예측하게 한 뒤, '
  '이 모델이 만든 "꿈" 안에서만 정책을 학습해도 실제 환경에서 동작함을 보였다. Dreamer 계열 [3]은 잠재 공간에서 상상한 궤적을 따라 가치와 정책을 '
  '역전파로 학습해 이를 연속 제어와 Atari로 넓혔고, MuZero [4]는 보상·가치·정책을 예측하는 학습된 모델 위에서 몬테카를로 트리 탐색을 수행해 '
  '규칙을 모르는 채로도 바둑과 Atari에서 계획이 가능함을 보였다.')
p('이 연구들에서 되풀이되는 어려움은 모델 오류의 누적이다. 모델의 예측을 다시 입력으로 넣어 여러 단계를 굴리면 작은 오류가 다음 입력을 바꾸고, '
  '그 위에서 또 오류가 생긴다. 그래서 실제 계획에서는 짧은 구간만 내다보고 첫 행동만 실행한 뒤 새 관측으로 다시 계획하는 receding horizon 방식이 '
  '널리 쓰이며, 본 연구의 에이전트도 같은 구조를 따른다. 또 하나 기억할 점은, 아무리 정확한 모델이 있어도 그 위에서 후보를 만들고 고르는 절차가 '
  '부실하면 성과가 나오지 않는다는 것이다. 월드 모델을 평가할 때 예측 정확도와 에이전트 성과를 따로 재야 하는 이유다.')

h2_('2.3 LLM 에이전트와 월드 모델')
p('LLM 에이전트는 보통 ReAct [5]처럼 관측과 이력을 프롬프트로 받아 다음 행동을 바로 생성한다. 명시적인 전이 모델은 없고, LLM이 사전학습에서 '
  '얻은 세상 지식이 암묵적인 모델 노릇을 한다. 이 방식은 간단하지만 행동의 결과를 확인하지 않고 행동하므로, 잠긴 문을 계속 열려 하거나 닫힌 상자 '
  '속 물건을 집으려 하는 식의 무효 행동을 되풀이한다. 이를 보완하려고 LLM 자체를 월드 모델로 쓰는 연구가 이어졌다. 표 1에 본 연구와 관련이 깊은 '
  '것들을 정리했다.', keep=True)
p('표 1. LLM을 월드 모델로 쓴 대표 연구', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['연구', '접근', '본 연구와의 연결'], [
    ('RAP [6]', 'LLM을 추론 에이전트와 월드 모델 양쪽으로 두고, 월드 모델이 예측한 다음 상태 위에서 MCTS로 추론·행동 경로를 탐색한다.',
     'LLM의 예측을 실행 전 후보 비교에 쓴 초기 사례'),
    ('WebDreamer [7]', '웹 에이전트가 후보 행동의 결과 페이지를 LLM으로 상상하고, 그 결과를 평가해 행동을 고른다.',
     '후보 → 결과 예측 → 평가 → 실행이라는 구조가 본 연구와 같다'),
    ('Wang 등 [8]', '텍스트 게임의 상태 전이 데이터로 LLM의 시뮬레이터 능력을 직접 측정했다. GPT-4도 상태 전이 정확도가 약 60%에 그쳤다.',
     '에이전트 성과와 별개로 전이 정확도를 재야 한다는 근거'),
    ('Xie 등 [9]', '행동의 실행 조건(precondition)과 효과(effect)를 예측하도록 LLM을 파인튜닝했다.',
     '"실행 여부 + 상태 변수" 질의와 환경 전이 학습(B·D)에 직접 대응'),
    ('Word2World [10]', '다섯 개 텍스트 환경에서 다음 상태 예측 모델을 학습하고, 데이터·모델 크기에 따른 변화와 에이전트 이득이 생기는 조건을 분석했다.',
     '학습한 월드 모델의 일반화 범위를 묻는 RQ4, 향후 연구의 비교 대상'),
], [28, 88, 68], size=9)
p('이 밖에 전이 규칙을 파이썬 코드로 써 내려가며 월드 모델을 만드는 WorldCoder [11]처럼 문장 대신 프로그램을 생성하는 접근도 있지만, '
  '표 1의 연구들은 모두 다음 상태를 글로 생성한다. 여기서 두 가지 쟁점이 생긴다.')
rich([('첫째는 생성된 상태의 일관성이다. ', True),
      ('다음 상태를 자유롭게 쓰게 하면 바뀐 사실을 추가하면서 바뀌기 전 사실을 지우지 않는 일이 생긴다. 사과를 집은 뒤에 "사과가 탁자 위에 있다"와 '
       '"사과가 인벤토리에 있다"가 함께 나오면 문장으로는 자연스러워도 하나의 상태로 쓸 수 없다. 이런 출력을 어떻게 처리하느냐가 에이전트 성과를 좌우한다.', False)])
rich([('둘째는 다단계 예측과 일반화다. ', True),
      ('자신의 예측을 다시 입력받는 자유 rollout에서는 오류가 누적된다. 또 환경 전이로 파인튜닝한 모델은 학습한 구조에서는 정확하지만, '
       '구조가 바뀌었을 때 얼마나 유지되는지는 따로 재 보기 전에는 알 수 없다. 새 이름에 대한 일반화와 새 구조에 대한 일반화는 다른 문제다.', False)])

h2_('2.4 결정 모델 JEV')
p('JEV는 TypeSafe AI가 2026년 9월에 공개한 System One Model이다 [12]. 이름처럼 긴 추론 대신 빠른 판단을 돌려주는 모델로, 자유 형식 문장을 '
  '생성하는 대신 상태(텍스트 또는 JSON)와 미리 정의한 질문을 받아 형식이 고정된 답과 확률을 돌려준다. 질문 유형은 세 가지다.', keep=True)
p('표 2. JEV의 질문 유형', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['유형', '반환 정보', '본 연구에서의 쓰임'], [
    ('Noul', '참(yes)일 확률', '예/아니오 판단. 보조적으로만 쓸 수 있다'),
    ('Choice', '선택된 값과 선택지별 확률', '실행 여부와 상태 변수 예측의 핵심 인터페이스'),
    ('Score', '순서형 점수와 수준별 확률', '사용하지 않음'),
], [24, 60, 100], size=9)
p('내부 구조와 학습 목적은 공개되어 있지 않고(RLCD라고만 소개된다), 사용자별 파인튜닝도 제공되지 않는다. 그래서 JEV는 공개된 확률 인터페이스만으로, '
  '이 과제의 추가 학습 없이(frozen) 쓴다. 요청 하나의 비용은 약 $0.00007이며, 본 실험에서는 episode 하나에 평균 129회를 요청해 약 $0.010이 들었다.')
p('본 연구의 착안은 단순하다. LLM에게 다음 상태를 통째로 쓰게 하는 대신, "이 행동 뒤 열쇠는 어디에 있는가? 상자 안 / 인벤토리 / 탁자 위 / …"처럼 '
  '상태 변수 하나하나를 선택지로 묻고, 그 답을 모아 다음 상태를 구성한다. 변수마다 값이 정확히 하나씩 나오므로 앞서 말한 일관성 문제가 '
  '구조적으로 사라지고, 확률이 함께 나오므로 불확실성을 계획에 반영할 수 있다. 이 질의 방식을 결정형, 다음 상태를 글로 쓰게 하는 방식을 생성형이라고 부른다.')

# ---- 3
h1_('3. 연구 질문과 비교 설계')
h2_('3.1 연구 질문')
bullets([
    ('RQ1. ', '월드 모델을 붙이면 에이전트의 과제 성공률이 실제로 오르는가? 실패한 행동을 기억하는 것만으로는 설명되지 않는 이득이 있는가?'),
    ('RQ2. ', '같은 LLM에서 결정형과 생성형 질의 방식은 성과와 형식 오류에 어떤 차이를 만드는가?'),
    ('RQ3. ', '환경 전이로 파인튜닝한 LLM은 과제 학습 없이 쓴 JEV와 비교해 어느 수준에 이르는가?'),
    ('RQ4. ', '여러 step을 굴렸을 때, 그리고 학습에 없던 구조에서 예측 정확도는 어떻게 변하는가?'),
])
p('RQ1–RQ3은 에이전트가 실제로 과제를 푸는 closed-loop 평가로, RQ4는 에이전트와 분리한 전이 예측 평가로 답한다.')

h2_('3.2 2×2 설계와 비교군')
p('비교의 뼈대는 질의 방식(결정형·생성형) × 학습 여부(없음·환경 전이 학습)의 2×2다. 여기에 JEV와 월드 모델 없는 기준선, '
  '환경의 정답을 쓰는 참조선을 더해 모두 아홉 개 비교군을 둔다.', keep=True)
p('표 3. 2×2 설계', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['', '학습 없음', '환경 전이로 학습'], [
    ('결정형 (선택지 확률)', 'A  JEV (frozen)    ·    B0  Qwen3-4B', 'B  Qwen3-4B + LoRA'),
    ('생성형 (다음 상태를 글로 씀)', 'D0  Qwen3-4B', 'D  Qwen3-4B + LoRA'),
    ('월드 모델 없음', 'C  정책만    ·    C_fm  정책 + 실패 기억', '—'),
], [48, 72, 64], size=9, center_cols=(1, 2))
p('표 4. 비교군 아홉 개', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['표기', '후보를 고르는 방법 / 월드 모델', '질의 방식', '과제 학습'], [
    ('C', '정책 1순위 계획의 첫 행동을 그대로 실행', '—', '—'),
    ('C_fm', '이 상태에서 실패한 적 없는 행동 중 정책 순위가 가장 높은 것', '—', '—'),
    ('validity', '엔진으로 실행 가능 여부만 확인한다 (결과는 모른다)', '정답 참조', '—'),
    ('oracle', '엔진으로 후보를 실제로 실행해 결과 상태를 본다', '정답 참조', '—'),
    ('A', 'JEV jev-1.13.0', '결정형', '없음'),
    ('B0', 'Qwen3-4B', '결정형', '없음'),
    ('D0', 'Qwen3-4B', '생성형', '없음'),
    ('B', 'Qwen3-4B + LoRA', '결정형', '전이 6,000개'),
    ('D', 'Qwen3-4B + LoRA', '생성형', '전이 6,000개 (B와 같은 데이터)'),
], [22, 96, 26, 40], size=9, center_cols=(2, 3))
p('C와 C_fm의 차이는 "실패한 행동을 반복하지 않는 것"만의 효과다. 월드 모델을 쓰는 비교군에도 같은 실패 기억이 들어 있으므로, 월드 모델의 이득은 '
  'C가 아니라 C_fm과 견주어야 한다. validity와 oracle의 차이는 "실행 여부를 넘어 결과까지 예측할 때 얻는 이득"이다. oracle은 같은 후보와 같은 '
  '계획기에 완벽한 예측을 넣은 것이므로 성능 상한이 아니라 진단 기준이다. 정책이 필요한 행동을 후보에 내지 않으면 oracle도 실패한다.')

h2_('3.3 통제한 것')
p('모든 비교군은 같은 정책(Qwen3-4B [14]), 같은 후보 생성 설정, 같은 계획기와 효용, 같은 실패 기억을 쓴다. 정책의 난수 seed는 (world, 시작 상태, step)으로만 '
  '정해지고 비교군 이름은 들어가지 않으므로, 같은 상태에서는 모든 비교군이 같은 후보를 받는다. LLM 월드 모델(B0·B·D0·D)은 정책과 같은 Qwen3-4B '
  '가중치를 쓰고, B·D는 여기에 LoRA [15] 어댑터를 얹는다. 정책이 후보를 만들 때는 어댑터를 끄고 월드 모델로 쓸 때만 켜므로, 어댑터 유무가 후보에 영향을 주지 않는다.')
p('B0와 D0, B와 D는 같은 모델이므로 둘의 차이는 질의 방식(과 그에 따라 생기는 확률·beam의 유무)에서 온다. 반면 A와 B·D의 비교는 범용 결정 모델과 '
  '이 과제의 전이로 학습한 모델의 비교이며, 모델 크기와 학습 이력이 함께 다른 시스템 비교로 읽어야 한다. JEV에 "학습 없음"이라고 쓴 것은 '
  '이 과제의 추가 학습을 하지 않았다는 뜻이다.')

# ---- 4
h1_('4. 과제 환경: TextWorld')
p('TextWorld [13]은 텍스트 어드벤처 게임을 생성하고 실행하는 연구용 환경이다. 본 연구는 TextWorld 1.7.0의 symbolic JSON backend를 써서 '
  '자연어 관측 대신 현재 상태의 사실 목록(예: in(key, box), locked(door))을 그대로 받는다. 과제는 방 두 개로 이루어진다(그림 1). 시작 방에는 '
  '열쇠가 든 닫힌 상자와 탁자가 있고, 탁자 위에는 목표 음식과 미끼 음식이 놓여 있다. 두 방 사이의 문은 잠겨 있고 상자 속 열쇠로만 열린다. '
  '목표는 목표 음식을 들고, 문을 연 채, 건너편 방에 있는 것이다. 최단 해법은 상자 열기 → 열쇠 집기 → 문 잠금 해제 → 목표 음식 집기 → 문 열기 → '
  '이동의 여섯 행동이며, episode는 최대 30 step이다.')
picture(QA / 'world.png', '그림 1. 과제 구성. 시작 방의 상자 속 열쇠로 잠긴 문을 열고, 목표 음식을 들고 건너편 방으로 간다.', width=150)
p('표 5. 과제의 성격', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['관측', '동역학', '목표', '성공 기준'], [
    ('현재 물리 상태의 사실 목록을 그대로 제공 (완전관측)', '결정적 전이. 실행되지 않는 명령은 아무것도 바꾸지 않는다',
     '목표 음식 소지 + 문 열림 + 건너편 방 도달 (목표 원자 3개)', '30 step 이내 목표 원자 3개 동시 충족'),
], [46, 50, 48, 40], size=9, align_first_center=False)
p('이 환경을 고른 이유는 세 가지다. 상태가 완전히 관측되고 전이가 결정적이므로 월드 모델의 예측을 엔진의 정답과 변수 단위로 비교할 수 있다. '
  '같은 엔진 정답을 oracle 비교군과 B·D의 학습 데이터로 쓸 수 있어, 월드 모델만 교체한 공정한 비교가 된다. 그리고 과제가 짧고 구조가 분명해서 '
  '실패의 원인을 "정책이 필요한 후보를 내지 않았다"와 "월드 모델이 잘못 예측했다"로 나눠 볼 수 있다.')
p('world는 같은 구조에 방과 물체의 이름만 다르게 생성했고, dev 12개(개발), train 150개(B·D 학습), val 20개(checkpoint 선택과 설정 보정), '
  'test 24개(최종 평가)로 나눴다. test의 이름은 학습과 보정에 쓰지 않은 단어다. 여기에 더해, 학습에 없던 구조(곁방, 미끼 열쇠, 음식이 든 둘째 용기)를 '
  '넣은 새 구조 X1을 전이 예측 평가에만 따로 쓴다.')

# ---- 5
h1_('5. 제안 방법: 결정 모델을 월드 모델로')
h2_('5.1 에이전트의 한 step')
picture(QA / 'architecture.png', '그림 2. 에이전트의 한 step. 월드 모델만 바꾸고 나머지는 모든 비교군이 공유한다.', width=170)
p('매 step 정책이 현재 상태와 목표, 명령 목록, 최근 시도 이력을 보고 길이 2의 후보 행동열을 여러 개 제안한다. 계획기는 각 후보의 앞부분(길이 1과 2)을 '
  '월드 모델로 굴려 결과 상태를 예측하고, 효용이 가장 높은 후보의 첫 행동 하나만 실제로 실행한다. 다음 step에는 새 관측에서 처음부터 다시 계획한다. '
  '예측이 틀려도 다음 step에서 실제 상태를 다시 받으므로 오류가 계속 끌려가지 않는다. 또 현재 상태에서 이미 실패한 명령으로 시작하는 후보는 '
  '평가에서 뺀다(실패 기억). 이 정보는 에이전트가 실제로 관측한 것이므로 모든 계획기 비교군에 똑같이 넣었다.')

h2_('5.2 한 단계 예측: 결정형')
p('월드 모델에는 언제나 "지금 상태에서 명령 하나"만 묻고, 여러 step은 예측한 상태를 다시 넣어 굴린다. 전이 규칙은 주지 않는다. 결정형은 행동 하나에 '
  '질문 일곱 개를 묻는다. 명령이 실행되는지와, 실행 뒤 상태 변수 여섯 개의 값이다.', keep=True)
p('표 6. 상태 변수와 선택지', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['질문', '선택지의 예'], [
    ('실행 여부', 'executes / fails'),
    ('플레이어 위치', '현재 방 / 연결된 다른 방'),
    ('열쇠·목표 음식·미끼 음식의 소재 (변수 3개)', '인벤토리 / 상자 안 / 탁자 위 / 방 바닥 / 소멸 등'),
    ('상자 상태, 문 상태 (변수 2개)', '열림 / 닫힘 / 잠김 등 해당 물체가 가질 수 있는 상태'),
], [70, 114], size=9, align_first_center=False)
p('각 변수의 가능한 값을 선택지로 주고, 돌아온 확률에서 가장 높은 값을 골라 다음 상태를 만든다. 바뀌지 않은 사실은 그대로 둔다.')
mathline(2)
p('qⱼ는 변수 j에 대한 선택지 확률, p는 실행 확률이다. JEV는 API의 Choice 질문을 그대로 쓰고, Qwen 결정형은 같은 질문을 "A. 상자 안 / B. 인벤토리 / …" '
  '형태로 넣은 뒤 다음 토큰 분포에서 선택지 글자의 확률만 골라 정규화한 값을 쓴다. 문장은 생성하지 않는다.')
p('표 7은 실제로 보내는 질문의 문구다. 상태는 물체의 id와 이름, 사실 목록을 담은 JSON으로 함께 준다.', keep=True)
p('표 7. 결정형 질의의 실제 문구', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['질문', '문구'], [
    ('상태 변수\n(변수마다 하나씩)', 'Starting from current_state, attempt exactly this one command: {명령}. Apply rollout_convention. Report the resulting state AFTER the command, not the current state. Question: {변수 질문, 예: Where is the brass key?}  Options: {가능한 값}'),
    ('실행 여부\n(별도 요청)', 'Consider the command below against current_state as it is now. Command: {명령}. Would this command actually execute, or would it fail because its requirements are not met in that state? Answer about whether it RUNS, not whether it is useful.  Options: executes / fails'),
], [30, 154], size=8.5)
p('실행 여부는 상태 질문과 따로 묻는다. 예비 실험에서 "명령을 실행했다고 치고 답하라"는 상태 질문들과 한 요청에 넣으면, 실행될 수 없는 명령의 실행 확률이 '
  '0.29에서 0.58로 올라가 실행을 전제하는 쪽으로 판단이 쏠렸다. 분리하자 행동 순서를 바꿨을 때 예측도 그에 맞게 바뀌는 비율이 18쌍 중 4쌍에서 14쌍으로 '
  '올랐다. 처음 설계는 행동열 전체 뒤의 상태를 한 번에 묻는 방식이었는데, 이 순서 검사를 전혀 통과하지 못해(0/18) 한 걸음씩 묻는 재귀 방식으로 바꿨다.')

h2_('5.3 확률 분기와 다단계 rollout')
p('실행 확률 p로 예측 상태에 도달하는 가지와, 1−p로 아무것도 바뀌지 않는 실패 가지를 나눠 가중치를 매긴다. 행동열을 따라 가지를 재귀적으로 이어 가되 '
  '폭 4의 beam만 유지하고 가중치 0.02 미만은 버린다. 끝 상태들의 가중 평균으로 목표 충족 확률과 진행도를 계산하고, 실패 가지의 질량을 더해 '
  '"실행되지 않을 명령 수"의 기댓값을 얻는다. 변수별 최빈값을 조합한 상태가 변수들의 결합 분포를 정확히 나타내는 것은 아니며, beam은 주로 '
  '실행·실패 경로의 불확실성을 전달한다.')

h2_('5.4 생성형 월드 모델과 파싱 규칙')
p('생성형(D0·D)은 같은 상태와 명령을 주고 "result: executes / fails" 한 줄과 결과 상태의 사실 목록을 쓰게 한 뒤 파서로 읽는다(greedy 디코딩, '
  '최대 320 토큰, 형식 예시 1개). 모든 상태 변수에 대해 값이 정확히 하나 있어야 하고, 없거나 둘 이상이면 파싱 실패다. 파싱 실패는 고정 seed로 한 번 '
  '다시 생성하고, 그래도 실패하면 "실행 안 됨, 상태 유지"로 처리한다. 생성형에는 확률이 없으므로 실행 여부는 0 또는 1이고 beam은 사실상 하나다.')
picture(EXAMPLE, '그림 3. 결정형과 생성형의 한 step 예측 예. 생성형이 사과를 두 곳에 두면 파싱 실패로 처리된다.', width=118)
p('그림 3은 두 방식의 차이를 한 예로 보여 준다. take apple from table 뒤에 생성형이 on(apple, table)과 in(apple, inventory)를 함께 쓰면 '
  '사과의 소재를 정할 수 없어 파싱 실패가 된다. 결정형은 사과의 소재를 한 질문으로 물어 값을 하나만 받는다.')

h2_('5.5 후보 평가')
p('모든 계획기 비교군은 같은 효용으로 후보를 고른다.', keep=True)
mathline(3)
p('conj는 예측한 끝 상태에서 목표 원자 3개(음식 소지, 문 열림, 건너편 방)가 모두 만족될 확률, progress는 만족된 원자의 비율, E[N_invalid]는 '
  '실행되지 않을 명령 수의 기댓값, h는 후보 길이, prior는 정책 순위에서 얻은 선호(1순위 1, 꼴찌 0)다. 계수는 실험 전에 정했고 한 번도 바꾸지 않았다. '
  'oracle과 validity는 월드 모델 대신 엔진 정답으로 같은 항을 계산한다. validity는 실행 여부만 알므로 conj와 progress를 0으로 둔다.')

# ---- 6
h1_('6. 실험 설정')
p('표 8. 최종 평가 설정', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['항목', '설정'], [
    ('데이터 분할', 'dev 12 / train 150 / val 20 / test 24 world'),
    ('학습 전이', 'train world에서 뽑은 엔진 정답 전이 6,000개 (B·D 동일)'),
    ('정책', 'Qwen3-4B bf16, thinking 끔, temperature 0.7, top_p 0.9. 한 step에 두 번 샘플링해 길이 2 후보 최대 16개'),
    ('계획과 종료', '후보 앞부분(길이 1·2)을 평가, 첫 행동만 실행. 30 step 이내 목표 충족 시 성공'),
    ('평가 규모', 'test 24 world × 시작 상태 4 × 정책 seed 3 = 비교군당 288 episode, 아홉 비교군 2,592 episode'),
    ('LoRA 학습', 'r=16, α=32, dropout 0.05, lr 1e-4, 2 epoch, seed 0. 정답 토큰에만 손실'),
    ('checkpoint 선택', 'val 전이 200개의 한 step 정확도가 가장 높은 epoch'),
    ('전이 예측 평가', '같은 구조 800 rollout, 새 구조 X1 400 rollout, 길이 4, 자유 rollout'),
    ('신뢰구간', 'world 단위로 짝지은 bootstrap 4,000회, 95%'),
], [36, 148], size=9)
p('B·D의 학습 데이터는 train world의 여러 상태에서 실행되는 명령과 실행되지 않는 명령을 반반에 가깝게 섞어 엔진으로 실행한 전이 6,000개다. '
  'B는 전이 하나를 질문 일곱 개(선택지 글자 하나가 정답)로, D는 "result" 줄과 사실 목록 하나로 배운다. JEV의 출력은 학습에도 checkpoint 선택에도 '
  '쓰지 않았다. JEV 이용약관이 출력으로 다른 모델을 학습시키는 것을 제한하기 때문이다. 학습 뒤 val 전이 200개에서의 한 step 정확도는 B 100%, D 99.5%였다'
  '(학습 전 Qwen은 52%, 40%).')
p('과제 난이도와 정책 프롬프트(용기와 열쇠에 관한 일반 규칙 등)는 예비 실험을 거쳐 dev·val world에서 월드 모델 없는 기준선(C, validity, oracle)의 '
  '성공률만 보고 정했다. 기준은 완벽한 예측을 넣은 oracle이 충분히 높게 풀 수 있을 것(90% 이상), oracle과 validity의 차이가 충분히 클 것(15%p 이상), '
  '후보 부족으로 인한 oracle 실패가 적을 것이었다. test world의 결과를 보고 설정이나 코드를 바꾼 일은 없다.')
p('성공률은 30 step 안에 목표 원자 3개를 모두 만족한 episode의 비율, 무효율은 실제로 실행한 행동 중 실행되지 않은 것의 비율이다. 전이 예측의 '
  '상태 완전일치는 깊이 k에서 beam 1위 상태의 변수가 모두 정답과 같은 비율이고, 생성형이 파싱에 실패한 step부터 그 rollout의 이후 깊이는 모두 틀린 '
  '것으로 센다. 바닥선으로 아무것도 바뀌지 않는다고 예측하는 persistence를 둔다. 같은 world의 episode는 서로 닮았으므로 신뢰구간은 world를 '
  '다시 뽑는 bootstrap으로 계산한다.')

# ---- 7
h1_('7. 결과')
h2_('7.1 과제 성공률')
p('표 9는 아홉 비교군의 성공률과 무효율, 월드 모델 비용이다. 대괄호는 world 단위 95% bootstrap 신뢰구간이다.', keep=True)
p('표 9. 과제 성공률, 무효율, episode당 월드 모델 비용', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['비교군', '성공률 % [95% CI]', '무효율 %', 'episode당 월드 모델 비용'], [
    ('C', '17.0 [7.3, 29.2]', '87.8', '—'),
    ('C_fm', '79.5 [70.5, 87.5]', '35.0', '—'),
    ('validity', '85.8 [77.8, 92.7]', '10.8', '—'),
    ('oracle', '92.7 [85.8, 97.9]', '12.8', '—'),
    ('A  JEV', '93.4 [85.1, 98.6]', '15.4', 'API 요청 129회, $0.010'),
    ('B0  결정형', '83.0 [73.3, 91.3]', '48.6', 'GPU 49초'),
    ('D0  생성형', '59.0 [48.6, 69.1]', '36.3', 'GPU 115초'),
    ('B  결정형 + 학습', '92.7 [85.8, 97.9]', '13.2', 'GPU 42초 (학습 13.9 GPU-h 별도)'),
    ('D  생성형 + 학습', '92.7 [85.8, 97.9]', '13.2', 'GPU 73초 (학습 4.5 GPU-h 별도)'),
], [40, 46, 26, 72], size=9, center_cols=(1, 2, 3))

h2_('7.2 월드 모델의 효과 (RQ1)')
p('월드 모델이 없는 C는 17.0%였고, 실패한 행동을 반복하지 않게만 해도(C_fm) 79.5%로 올랐다. 실패 기억의 효과가 그만큼 크다. 이를 통제한 뒤에도 '
  'JEV(A)는 C_fm보다 13.9%p [2.1, 24.3], 학습한 B·D는 13.2%p [2.8, 22.9] 높았다. 짝지은 288 episode 가운데 A가 C_fm보다 잘한 것이 50개, 못한 것이 '
  '10개였다. 실행 여부만 정확히 아는 validity(85.8%)와 결과까지 아는 oracle(92.7%)의 차이 6.9%p [1.0, 13.9]는 "실행되는지"를 넘어 '
  '"무엇이 바뀌는지"를 예측할 때 추가로 얻는 이득이다.')

h2_('7.3 학습 없는 LLM: 결정형과 생성형 (RQ2)')
p('학습 없는 Qwen3-4B는 질문을 선택지로 바꾸기만 해도 59.0%(D0)에서 83.0%(B0)로 24.0%p [11.5, 35.8] 올랐다. D0는 실패 기억만 쓴 C_fm보다도 '
  '20.5%p 낮았다. 원인은 예측 내용이 아니라 출력 형식이다. D0는 한 step 예측의 26.0%에서 그림 3과 같이 파싱할 수 없는 상태를 냈고, 규칙에 따라 '
  '"실행 안 됨"으로 처리됐다. 음식을 집는 행동이 이렇게 처리되면 에이전트는 음식 집기를 피하게 되고, 결국 음식 없이 건너편 방에 들어간다. 이 상태에 '
  '빠진 episode는 D0 40%, C_fm 21%였고, 들어간 뒤 성공한 경우는 거의 없었다. 파싱에 성공한 응답만 보면 D0의 정확도는 B0보다 낮지 않았다. '
  '즉 결정형의 이점은 더 정확히 예측해서가 아니라 형식이 깨지지 않는 데서 생긴다.')
p('전이 예측 정확도에서는 JEV와 학습 없는 LLM의 차이가 컸다. 그림 4(a)에서 JEV는 4 step 뒤에도 상태의 77.9%를 정확히 맞혔지만, B0와 D0는 '
  '9.1%, 6.2%로 아무것도 바뀌지 않는다고 예측하는 persistence(15.6%)보다 낮았다. 그런데도 B0의 성공률(83.0%)은 C_fm과 큰 차이가 없었고 '
  'A(93.4%)와도 통계적으로 구분되지 않았다(+10.4%p [−1.7, 21.5]). 예측이 부정확해도 실패 기억과 매 step의 재계획이 있어, closed-loop 성공률은 '
  '예측 정확도만큼 벌어지지 않는다.')

h2_('7.4 JEV와 oracle, 파인튜닝한 LLM (RQ3)')
p('과제 학습 없이 쓴 JEV는 93.4%로 oracle(92.7%)과 같은 수준이었다(차이 +0.7%p [−1.7, 3.1]). 두 비교군의 성공 여부가 다른 episode는 288개 중 '
  '8개뿐이다. oracle이 실패한 21개 episode는 정책이 필요한 행동을 후보에 아예 내지 않은 경우 6개와, 후보에는 있었지만 길이 2의 근시안적 효용이 '
  '고르지 못한 경우 15개로 나뉜다. 둘 다 월드 모델이 고칠 수 없는 실패이므로, 이 과제에서 JEV는 월드 모델로서 할 수 있는 몫을 거의 다 한 셈이다.')
p('파인튜닝한 B·D는 288 episode 모두에서 oracle과 성공 여부가 같았고, 행동 순서까지 같은 episode가 282개와 284개였다. 파싱 실패도 0.1%로 줄어 '
  '성공률에서는 결정형과 생성형의 차이가 사라졌다. 결국 좋은 월드 모델 셋(A, B, D)은 closed-loop 성공률로는 구분되지 않으며, 이들의 차이는 '
  '다음의 전이 예측 평가에서 드러난다.')

h2_('7.5 다단계 예측과 새 구조 (RQ4)')
picture(ROLLOUT, '그림 4. 자유 rollout의 깊이별 상태 완전일치율. (a) 같은 구조의 test world, (b) 학습에 없던 구조 X1. 음영은 95% CI.', width=150)
p('표 10. 깊이 1과 4에서의 상태 완전일치율 (%)', size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, keep=True)
tbl(['모델', '같은 구조 k=1', '같은 구조 k=4', '새 구조 X1 k=1', '새 구조 X1 k=4'], [
    ('persistence (바닥선)', '35.0', '15.6', '40.3', '18.0'),
    ('A  JEV', '90.9', '77.9', '94.8', '85.5'),
    ('B0  결정형', '39.4', '9.1', '34.3', '4.3'),
    ('D0  생성형', '40.0', '6.2', '20.3', '2.3'),
    ('B  결정형 + 학습', '100.0', '98.5', '89.8', '71.5'),
    ('D  생성형 + 학습', '99.8', '96.8', '91.5', '66.5'),
], [48, 34, 34, 34, 34], size=9, center_cols=(1, 2, 3, 4))
p('같은 구조(a)에서 파인튜닝한 B·D는 4 step 뒤에도 98.5%, 96.8%로 JEV(77.9%)보다 정확했다. 다만 test world는 이름만 새롭고 구조는 학습과 같다. '
  '이름을 id로 바꾸면 test 전이의 69.5%가 학습 데이터에 그대로 있다. 곁방과 미끼 열쇠, 둘째 용기를 더한 새 구조 X1(b)에서는 순서가 뒤집혔다. '
  '4 step 정확도는 JEV 85.5%, B 71.5%, D 66.5%였고, 같은 구조에서 새 구조로 옮겼을 때의 하락폭은 B가 JEV보다 25.6%p [18.8, 33.9], '
  'D가 26.7%p [21.7, 31.6] 컸다. 파인튜닝의 이득은 학습한 구조 안에 머물렀다. 학습 없는 B0·D0는 새 구조에서도 persistence보다 낮았고, '
  '파인튜닝한 D의 파싱 실패는 같은 구조에서 0.1%, 새 구조에서 2.1%였다.')
p('JEV의 정확도가 새 구조에서 오히려 높게 나온 것은 X1의 행동열에 실행되지 않아 상태가 유지되는 step이 더 많기 때문이다. 두 세트의 수치를 '
  '직접 견주기보다 각 세트 안에서 모델 사이의 차이를 보는 것이 맞다.')

# ---- 8
h1_('8. 해석과 한계')
h2_('8.1 무엇을 말할 수 있는가')
bullets([
    '결정 모델 JEV는 과제 학습 없이도 LLM 에이전트의 월드 모델로 쓸 수 있었고, 성공률은 완벽한 예측을 넣은 기준과 같은 수준이었다.',
    'LLM을 학습 없이 쓸 때는 선택지로 묻는 결정형이 형식 오류를 막아 생성형보다 훨씬 나았다. 이 이점은 예측 내용이 아니라 형식의 견고성에서 왔다.',
    '환경 전이로 파인튜닝한 LLM은 학습한 구조에서는 가장 정확했지만 새 구조에서는 JEV보다 크게 떨어졌다. 파인튜닝은 "이 구조"를 배웠고, '
    'JEV는 구조를 배우지 않은 채로 쓸 만한 예측을 냈다.',
], numbered=True)

h2_('8.2 범위와 한계')
bullets([
    '과제 구조 하나와 LLM 하나(Qwen3-4B), 학습 seed 하나에서 얻은 결과다. 다른 환경이나 더 큰 모델로 일반화하려면 추가 실험이 필요하다.',
    '결정형과 생성형의 성공률 차이(24.0%p)의 크기는 파싱 실패를 "실행 안 됨"으로 처리하는 규칙에 달려 있다. 다른 처리 규칙에서는 차이가 달라질 수 있다.',
    '결정형의 이점이 질문 형식 자체에서 오는지, 함께 들어오는 확률과 beam에서 오는지는 분리하지 못했다. 상태 변수 틀을 연구자가 설계해야 한다는 비용도 있다.',
    'JEV의 내부 구조와 학습 목적이 공개되어 있지 않으므로 성능의 원인을 모델 내부에서 찾을 수 없고, 모델 버전이 바뀌면 같은 수치를 재현하기 어렵다. '
    '원시 요청과 응답은 모두 보관했다.',
    'closed-loop 계획 길이는 2이며, 긴 horizon 계획이나 부분 관측 환경은 다루지 않았다.',
])

h2_('8.3 진행 상황과 향후 계획')
p('이 결과로 KIIS 2026 추계학술대회 구두발표 논문(A4 2쪽)을 투고했다. 10월 2일 요약문이 접수됐고, 10월 23일에 최종 원고를 제출한다. 다음 단계로는 '
  'ScienceWorld·ALFWorld처럼 더 어렵고 부분 관측이 섞인 환경에서 공개된 LLM 월드 모델(Word2World)과 JEV를 같은 계획기 아래 비교하는 것, '
  '파싱 실패 처리 규칙과 스키마 제약 생성(constrained decoding)을 바꿔 가며 결정형 이점의 출처를 가르는 것, 학습 seed와 모델 크기를 늘려 파인튜닝의 '
  '일반화 범위를 더 넓게 재는 것을 계획하고 있다.')

# ---- 참고문헌
h1_('참고문헌')
refs = [
    'R. S. Sutton, "Dyna, an integrated architecture for learning, planning, and reacting," SIGART Bulletin, vol. 2, no. 4, pp. 160–163, 1991.',
    'D. Ha and J. Schmidhuber, "World Models," arXiv:1803.10122, 2018.',
    'D. Hafner, T. Lillicrap, J. Ba, and M. Norouzi, "Dream to Control: Learning Behaviors by Latent Imagination," Proc. ICLR, 2020.',
    'J. Schrittwieser et al., "Mastering Atari, Go, chess and shogi by planning with a learned model," Nature, vol. 588, pp. 604–609, 2020.',
    'S. Yao et al., "ReAct: Synergizing Reasoning and Acting in Language Models," Proc. ICLR, 2023.',
    'S. Hao et al., "Reasoning with Language Model is Planning with World Model," Proc. EMNLP, pp. 8154–8173, 2023.',
    'Y. Gu et al., "Is Your LLM Secretly a World Model of the Internet? Model-Based Planning for Web Agents," Transactions on Machine Learning Research, 2025.',
    'R. Wang et al., "Can Language Models Serve as Text-Based World Simulators?," Proc. ACL (Short Papers), pp. 1–17, 2024.',
    'K. Xie et al., "Making Large Language Models into World Models with Precondition and Effect Knowledge," Proc. COLING, pp. 7532–7545, 2025.',
    'Y. Li et al., "From Word to World: Can Large Language Models be Implicit Text-based World Models?," arXiv:2512.18832, 2025.',
    'H. Tang, D. Key, and K. Ellis, "WorldCoder, a Model-Based LLM Agent: Building World Models by Writing Code and Interacting with the Environment," Proc. NeurIPS, 2024.',
    'TypeSafe AI, "Introducing System One Models & Jev," https://typesafe.ai/blog/introducing-system-one-models-and-jev (열람 2026년 10월 2일).',
    'M.-A. Côté et al., "TextWorld: A Learning Environment for Text-based Games," Proc. IJCAI Computer Games Workshop, 2018.',
    'A. Yang et al., "Qwen3 Technical Report," arXiv:2505.09388, 2025.',
    'E. J. Hu et al., "LoRA: Low-Rank Adaptation of Large Language Models," Proc. ICLR, 2022.',
]
for i, t in enumerate(refs, 1):
    z = doc.add_paragraph(); z.paragraph_format.space_after = Pt(2); z.paragraph_format.line_spacing = Pt(12.5)
    z.paragraph_format.left_indent = Mm(7); z.paragraph_format.first_line_indent = Mm(-7)
    z.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = z.add_run(f'[{i}] {t}'); r.font.size = Pt(8.5)

doc.save(OUT)
print('saved', OUT)
